"""
Splits MASTER SHEET.xlsx (multiple runs in one file, Format C)
into individual Excel files (Format A) in test_data/before/.
"""
import openpyxl
from openpyxl import Workbook
from pathlib import Path

SRC = Path("MASTER SHEET.xlsx")
OUT = Path("test_data/before")
OUT.mkdir(parents=True, exist_ok=True)

RUN_DEFS = [
    (1,  "RUN_01"), (7,  "RUN_03"), (10, "RUN_04"),
    (13, "RUN_05"), (16, "RUN_06"), (19, "RUN_07"),
    (22, "RUN_08"), (25, "RUN_09"), (28, "RUN_10"),
]

wb_src = openpyxl.load_workbook(SRC)
ws_src = wb_src.active

for col, run_name in RUN_DEFS:
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = run_name
    ws_out.append(["Peak Position (cm-1)", "Peak Intensity"])

    written = 0
    for row in ws_src.iter_rows(min_row=3, min_col=col, max_col=col + 1, values_only=True):
        if row[0] is not None and isinstance(row[0], (int, float)):
            ws_out.append([float(row[0]), float(row[1]) if row[1] is not None else 0.0])
            written += 1

    out_path = OUT / f"{run_name}.xlsx"
    wb_out.save(out_path)
    print(f"  {run_name}.xlsx  →  {written} peaks")

print(f"\nDone. {len(RUN_DEFS)} files written to {OUT}/")
